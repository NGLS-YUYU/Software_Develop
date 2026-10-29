"""Prepare isolated E2E fixtures in uams_test (run with backend venv).

Creates a temporary course/task/class/student for grade import, writes
student/teacher/course/grade xlsx files, and a metadata JSON.
"""
import json
import urllib.request
import urllib.error
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

API = "http://127.0.0.1:8010/api/v1"
OUT = Path("/tmp/uams-e2e")
OUT.mkdir(exist_ok=True)
META = OUT / "meta.json"


def api(path, tok=None, method="GET", body=None, raw=False):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(
        API + path, data=data, method=method,
        headers={"Content-Type": "application/json"} if data else {},
    )
    if tok:
        req.add_header("Authorization", "Bearer " + tok)
    try:
        with urllib.request.urlopen(req, timeout=20) as x:
            payload = x.read()
            return x.status, payload if raw else (json.loads(payload) if payload else None)
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


st, admin = api("/auth/login", method="POST", body={
    "username": "admin", "real_name": "admin", "password": "admin123"})
assert st == 200, (st, admin)
tok = admin["access_token"]

st, colleges = api("/colleges?page_size=5", tok)
college = colleges["items"][0]
st, majors = api(f"/majors?page_size=50&college_id={college['id']}", tok)
major = majors["items"][0]
st, classes = api(f"/classes?page_size=50&major_id={major['id']}", tok)
klass = classes["items"][0]
st, teachers = api("/teachers?keyword=T001", tok)
teacher = teachers["items"][0]
st, term = api("/terms/active", tok)
semester = f"{term['academic_year']}-{term['semester']}"

# --- temporary grade-import teaching class ---
st, course = api("/courses", tok, "POST", {
    "code": "E2EACC-GRADE", "name": "E2E成绩导入课程", "college_id": college["id"],
    "course_type": "ELECTIVE", "credits": 2, "total_hours": 32,
    "theory_hours": 32, "practice_hours": 0, "exam_type": "EXAM", "status": "ACTIVE",
})
if st == 409:
    st, page = api("/courses?keyword=E2EACC-GRADE", tok)
    course = page["items"][0]
assert st in (201, 200), (st, course)

st, task = api("/teaching-tasks", tok, "POST", {
    "semester": semester, "course_id": course["id"], "college_id": college["id"]})
if st == 409:
    st, page = api(f"/teaching-tasks?semester={semester}&course_id={course['id']}", tok)
    task = page["items"][0]
assert st in (201, 200), (st, task)

st, tclass = api("/teaching-classes", tok, "POST", {
    "code": "E2EACC-GRADE-01", "task_id": task["id"], "term_id": term["id"],
    "teacher_id": teacher["id"], "capacity": 10})
if st == 409:
    st, page = api(f"/teaching-classes?term_id={term['id']}&course_id={course['id']}", tok)
    tclass = page["items"][0]
assert st in (201, 200), (st, tclass)

# 添加排课后才能把学生加入教学班
st, rooms = api("/classrooms?page_size=5", tok)
room = rooms["items"][0]
scheduled = False
for day in range(7, 0, -1):
    st, _ = api(f"/teaching-classes/{tclass['id']}/schedules", tok, "POST", {
        "classroom_id": room["id"], "day_of_week": day,
        "start_period": 13, "end_period": 14, "start_week": 1,
        "end_week": 16, "week_type": "ALL"})
    if st in (200, 201):
        scheduled = True
        break
    if st == 409 and "已存在" not in str(_):
        # 冲突：换一天再试
        continue
assert scheduled, "无法为临时教学班排课"

st, student = api("/students", tok, "POST", {
    "student_no": "E2EACC_G1", "real_name": "成绩导入测试", "class_id": klass["id"],
    "college_id": college["id"], "major_id": major["id"], "enrollment_year": 2026,
})
if st == 409:
    st, page = api("/students?keyword=E2EACC_G1", tok)
    student = page["items"][0]
assert st in (201, 200), (st, student)

st, add = api(f"/teaching-classes/{tclass['id']}/students", tok, "POST",
              {"student_id": student["id"]})
if st not in (200, 201):
    st2, roster = api(f"/teaching-classes/{tclass['id']}/students", tok)
    assert st2 == 200 and any(r["student_id"] == student["id"] for r in roster), (st, add)

# --- xlsx fixtures ---
def xlsx(path, headers, row):
    wb = Workbook(); ws = wb.active; ws.append(headers); ws.append(row)
    wb.save(path)


xlsx(OUT / "students-ok.xlsx",
     ["学号", "姓名", "性别", "学院", "专业", "行政班", "入学年份", "邮箱", "手机"],
     ["E2EACC_S1", "导入学生甲", "女", college["name"], major["name"], klass["name"], 2026, "", ""])
xlsx(OUT / "students-bad.xlsx",
     ["学号", "姓名", "性别", "学院", "专业", "行政班"],
     ["E2EACC_S2", "导入学生乙", "未知X", college["name"], major["name"], klass["name"]])

xlsx(OUT / "teachers-ok.xlsx",
     ["工号", "姓名", "学院", "职称", "性别", "入职日期", "邮箱", "联系电话"],
     ["E2EACC_T1", "导入教师甲", college["name"], "讲师", "男", "2024-09-01", "", ""])
xlsx(OUT / "teachers-bad.xlsx",
     ["工号", "姓名", "学院", "性别"],
     ["E2EACC_T2", "导入教师乙", college["name"], "未知X"])

xlsx(OUT / "courses-ok.xlsx",
     ["课程代码", "课程名称", "学院", "学分", "总学时", "理论学时", "实践学时", "课程性质", "考核形式", "课程简介", "状态"],
     ["E2EACC_C1", "导入课程甲", college["name"], 2, 32, 32, 0, "专业选修", "考查", "", "启用"])
xlsx(OUT / "courses-bad.xlsx",
     ["课程代码", "课程名称", "学院", "学分", "总学时", "理论学时", "实践学时", "课程性质"],
     ["E2EACC_C2", "导入课程乙", college["name"], 2, 32, 30, 4, "专业选修"])

# grade template generated by backend, then filled for the temp student
st, blob = api(f"/imports/templates/grades?teaching_class_id={tclass['id']}", tok, raw=True)
assert st == 200, st
tpl = OUT / "grades-template.xlsx"
tpl.write_bytes(blob)
for suffix, regular, final in (("ok", 85, 90), ("bad", 150, 90)):
    wb = load_workbook(tpl)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    row_idx = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "E2EACC_G1":
            row_idx = r
            break
    assert row_idx, "temp student not in grade template"
    ws.cell(row_idx, headers.index("平时成绩") + 1).value = regular
    ws.cell(row_idx, headers.index("期末成绩") + 1).value = final
    wb.save(OUT / f"grades-{suffix}.xlsx")

# --- 只读：取受保护记录（有选课的学生 / 有教学任务的课程 / 有成绩的名单） ---
import re
from sqlalchemy import create_engine, text

BACKEND_ENV = Path(__file__).resolve().parents[3] / "backend" / ".env"
base_url = re.search(r"^DATABASE_URL=(.+)$",
                     BACKEND_ENV.read_text(encoding="utf-8"), re.M).group(1).strip()
test_url = base_url.rsplit("/", 1)[0] + "/uams_test?charset=utf8mb4"
engine = create_engine(test_url)
with engine.connect() as conn:
    row = conn.execute(text(
        "SELECT s.student_no FROM students s JOIN course_selections cs "
        "ON cs.student_id=s.id LIMIT 1")).first()
    protected_student_no = row[0]
    row = conn.execute(text(
        "SELECT c.code FROM courses c JOIN teaching_tasks t ON t.course_id=c.id LIMIT 1")).first()
    protected_course_code = row[0]
    row = conn.execute(text(
        "SELECT tc.code, s.student_no, tc.id FROM grades g "
        "JOIN teaching_classes tc ON tc.id=g.teaching_class_id "
        "JOIN students s ON s.id=g.student_id LIMIT 1")).first()
    graded_class = {"id": row[2], "code": row[0], "student_no": row[1]}
engine.dispose()

meta = {
    "college": {"id": college["id"], "name": college["name"]},
    "major": {"id": major["id"], "name": major["name"]},
    "class": {"id": klass["id"], "name": klass["name"]},
    "majors": [m["name"] for m in majors["items"] if m["college_id"] == college["id"]],
    "classes": [k["name"] for k in classes["items"] if k["major_id"] == major["id"]],
    "protected_student_no": protected_student_no,
    "protected_course_code": protected_course_code,
    "graded_class": graded_class,
    "term": {"id": term["id"], "semester": semester},
    "course": {"id": course["id"], "code": course["code"]},
    "task_id": task["id"],
    "tclass": {"id": tclass["id"], "code": tclass["code"]},
    "student": {"id": student["id"], "student_no": "E2EACC_G1"},
    "files": {k: str(OUT / f"{k}.xlsx") for k in
              ("students-ok", "students-bad", "teachers-ok", "teachers-bad",
               "courses-ok", "courses-bad", "grades-ok", "grades-bad")},
}
META.write_text(json.dumps(meta, ensure_ascii=False, indent=2))
print("PREP_OK", json.dumps({k: meta[k] for k in ("college", "major", "class", "tclass", "student")}, ensure_ascii=False))
