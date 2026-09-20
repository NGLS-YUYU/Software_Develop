"""FIX-005 真实 E2E：真实 HTTP + 真实 Excel + 真实 MySQL。

用法（先启动后端）：
    .venv/bin/python -m uvicorn app.main:app --port 8077
    .venv/bin/python scripts/e2e_fix005_import.py

流程与学生/教师/课程/成绩四种导入完全一致：下载模板 → 填表 → 上传 → 查库。
"""

import io
import sys
from decimal import Decimal
from pathlib import Path

import httpx
from openpyxl import load_workbook
from sqlalchemy import select

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.db.session import SessionLocal  # noqa: E402
from app.models.course import Course  # noqa: E402
from app.models.grade import Grade  # noqa: E402
from app.models.person import Student, Teacher  # noqa: E402
from app.models.user import User  # noqa: E402

BASE = "http://127.0.0.1:8077/api/v1"
P = "E2E5"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OK, FAIL = "✅", "❌"
results: list[tuple[bool, str]] = []


def check(cond: bool, label: str, detail: str = "") -> None:
    results.append((bool(cond), label))
    print(f"  {OK if cond else FAIL} {label}{(' — ' + detail) if detail else ''}")


def login(client: httpx.Client, username: str, password: str) -> dict:
    r = client.post(f"{BASE}/auth/login", json={"username": username, "password": password})
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def download_template(client, headers, kind: str, **params) -> bytes:
    r = client.get(f"{BASE}/imports/templates/{kind}", params=params, headers=headers)
    r.raise_for_status()
    return r.content


def sheets(template: bytes) -> list[str]:
    return load_workbook(io.BytesIO(template)).sheetnames


def headers_of(template: bytes) -> list[str]:
    ws = load_workbook(io.BytesIO(template)).worksheets[0]
    return [("" if c.value is None else str(c.value).strip()) for c in ws[1]]


def build(headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fill(template: bytes, values: dict[str, dict[str, object]]) -> bytes:
    """按学号把成绩写进系统模板（模拟教师填表）。"""
    wb = load_workbook(io.BytesIO(template))
    ws = wb.worksheets[0]
    hdr = headers_of(template)
    for r in range(2, ws.max_row + 1):
        no = str(ws.cell(r, hdr.index("学号") + 1).value or "").strip()
        for label, value in (values.get(no) or {}).items():
            ws.cell(r, hdr.index(label) + 1).value = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, headers, kind: str, data: bytes, name="data.xlsx"):
    return client.post(
        f"{BASE}/imports/{kind}", files={"file": (name, data, XLSX)}, headers=headers
    )


def cleanup() -> None:
    from sqlalchemy import delete

    from app.models.role import UserRole
    from app.models.selection import CourseSelection
    from app.models.teaching import Schedule, TeachingClass, TeachingTask

    with SessionLocal() as db:
        users = select(User.id).where(User.username.like(f"{P}%"))
        courses = select(Course.id).where(Course.code.like(f"{P}%"))
        tcs = select(TeachingClass.id).where(TeachingClass.code.like(f"{P}%"))
        # 顺序必须满足外键：成绩/选课/排课 → 教学班 → 教学任务 → 人员 → 课程
        db.execute(delete(UserRole).where(UserRole.user_id.in_(users)))
        db.execute(delete(Grade).where(Grade.student_id.in_(
            select(Student.id).where(Student.student_no.like(f"{P}%"))
        )))
        db.execute(delete(Grade).where(Grade.teaching_class_id.in_(tcs)))
        db.execute(delete(CourseSelection).where(CourseSelection.teaching_class_id.in_(tcs)))
        db.execute(delete(Schedule).where(Schedule.teaching_class_id.in_(tcs)))
        db.execute(delete(TeachingClass).where(TeachingClass.code.like(f"{P}%")))
        db.execute(delete(TeachingTask).where(TeachingTask.course_id.in_(courses)))
        db.execute(delete(Student).where(Student.student_no.like(f"{P}%")))
        db.execute(delete(Teacher).where(Teacher.teacher_no.like(f"{P}%")))
        db.execute(delete(User).where(User.username.like(f"{P}%")))
        db.execute(delete(Course).where(Course.code.like(f"{P}%")))
        db.commit()


def main() -> int:
    cleanup()
    with httpx.Client(timeout=30) as client:
        admin = login(client, "admin", "admin123")

        # ---------- 1. 教师批量导入 ----------
        print("\n[1] 教师批量导入")
        tpl = download_template(client, admin, "teachers")
        check(sheets(tpl) == ["导入模板", "填写说明"], "模板含「填写说明」工作表")
        check(headers_of(tpl)[:3] == ["工号", "姓名", "学院"], "模板首行是中文字段名")
        data = build(
            headers_of(tpl),
            [
                [f"{P}T1", "E2E教师甲", "计算机科学与技术学院", "副教授", "女", "13900000001"],
                [f"{P}T2", "E2E教师乙", "计算机科学与技术学院", "讲师", "男", ""],
            ],
        )
        r = upload(client, admin, "teachers", data)
        body = r.json()
        check(r.status_code == 200 and body["success"] == 2, "上传 2 名教师", str(body))
        with SessionLocal() as db:
            t = db.execute(select(Teacher).where(Teacher.teacher_no == f"{P}T1")).scalar_one()
            check(t.title == "ASSOCIATE_PROFESSOR" and t.gender == "FEMALE",
                  "MySQL：职称/性别中文已翻译为枚举", f"{t.title}/{t.gender}")
            check(t.user.real_name == "E2E教师甲", "MySQL：用户档案已建立")
        check(client.post(f"{BASE}/auth/login",
                          json={"username": f"{P}T1", "password": f"{P}T1"}).status_code == 200,
              "导入的教师可用「工号=初始密码」登录")

        # ---------- 2. 学生批量导入 ----------
        print("\n[2] 学生批量导入")
        tpl = download_template(client, admin, "students")
        data = build(
            headers_of(tpl),
            [
                [f"{P}S1", "E2E学生甲", f"{P}S1", "计算机科学与技术学院", "计算机科学与技术", "计算机科学与技术2026-1班"],
                [f"{P}S2", "E2E学生乙", f"{P}S2", "计算机科学与技术学院", "计算机科学与技术", "计算机科学与技术2026-1班"],
            ],
        )
        body = upload(client, admin, "students", data).json()
        check(body["success"] == 2, "上传 2 名学生", str(body))
        with SessionLocal() as db:
            s = db.execute(select(Student).where(Student.student_no == f"{P}S1")).scalar_one()
            check(s.user.real_name == "E2E学生甲", "MySQL：学生+用户已建立")

        # 错误行：学院不存在（部分成功策略）
        data = build(headers_of(tpl), [
            [f"{P}S3", "E2E学生丙", f"{P}S3", "不存在的学院", "计算机科学与技术", "计算机科学与技术2026-1班"],
        ])
        body = upload(client, admin, "students", data).json()
        check(body["failed"] == 1 and "学院不存在" in body["errors"][0]["message"],
              "错误行返回明细（行号+原因）", str(body["errors"][0]))

        # ---------- 3. 课程批量导入 ----------
        print("\n[3] 课程批量导入")
        tpl = download_template(client, admin, "courses")
        data = build(headers_of(tpl), [
            [f"{P}C1", "E2E导入课", "计算机科学与技术学院", 3.0, 48, 40, 8, "专业必修"],
        ])
        body = upload(client, admin, "courses", data).json()
        check(body["success"] == 1, "导入合法课程（48h = 40 + 8）", str(body))
        with SessionLocal() as db:
            c = db.execute(select(Course).where(Course.code == f"{P}C1")).scalar_one()
            check(c.assessment_type == "48" and c.practice_hours == 8,
                  "MySQL：考核方式由总学时推导，实践学时已落库",
                  f"{c.assessment_type}/{c.practice_hours}")

        data = build(headers_of(tpl), [
            [f"{P}C2", "E2E错学时课", "计算机科学与技术学院", 3.0, 32, 20, 8, "专业必修"],
        ])
        body = upload(client, admin, "courses", data).json()
        with SessionLocal() as db:
            exists = db.execute(select(Course).where(Course.code == f"{P}C2")).scalar_one_or_none()
        check(body["failed"] == 1 and exists is None,
              "理论+实践≠总学时 → 拒绝且未落库", str(body["errors"][0]))

        # ---------- 4. 成绩：模板动态生成 ----------
        print("\n[4] 成绩模板动态生成（按教学班 + 考核结构）")
        t1 = login(client, f"{P}T1", f"{P}T1")
        with SessionLocal() as db:
            teacher = db.execute(select(Teacher).where(Teacher.teacher_no == f"{P}T1")).scalar_one()
            c48p = db.execute(select(Course).where(Course.code == f"{P}C1")).scalar_one()
            # 直接建教学班（FIX-005 不负责开课；这里为 E2E 造一个真实教学班）
            from app.models.role import Role, RoleCode
            from app.models.teaching import TeachingClass, TeachingTask
            from app.models.term import AcademicTerm
            from app.models.selection import CourseSelection, SelectionState

            term = db.execute(select(AcademicTerm).where(
                AcademicTerm.academic_year == "2026-2027", AcademicTerm.semester == 1
            )).scalar_one()
            task = TeachingTask(semester="2026-2027-1", course_id=c48p.id,
                                college_id=c48p.college_id, status="CONFIRMED")
            db.add(task)
            db.flush()
            tc = TeachingClass(code=f"{P}TC1", task_id=task.id, course_id=c48p.id,
                               semester="2026-2027-1", teacher_id=teacher.id,
                               capacity=60, term_id=term.id, selection_status="OPEN")
            db.add(tc)
            db.flush()
            s1 = db.execute(select(Student).where(Student.student_no == f"{P}S1")).scalar_one()
            db.add(CourseSelection(student_id=s1.id, teaching_class_id=tc.id,
                                   course_id=c48p.id, semester="2026-2027-1",
                                   status=SelectionState.SELECTED))
            db.commit()
            tc_id, s1_id = tc.id, s1.id
            tc_code = tc.code
            _ = Role, RoleCode

        tpl = download_template(client, t1, "grades", teaching_class_id=tc_id)
        hdr = headers_of(tpl)
        check(hdr[:6] == ["学号", "姓名", "班级", "课程代码", "课程名称", "教学班"],
              "模板固定列由系统生成", str(hdr[:6]))
        check(hdr[6:] == ["平时成绩", "期中理论", "期中实践", "期末理论",
                          "期末实践", "理论补考", "实践补考"],
              "48h+实践 → 三段式理论/实践列（无「期中成绩」合并列）", str(hdr[6:]))
        ws = load_workbook(io.BytesIO(tpl)).worksheets[0]
        check(str(ws.cell(2, 1).value) == f"{P}S1" and str(ws.cell(2, 2).value) == "E2E学生甲",
              "模板已自动带出学生名单（老师不用填学号/姓名）")

        # ---------- 5. 成绩导入 + ScoreService 计算 ----------
        print("\n[5] 成绩导入（填表→上传→查库）")
        data = fill(tpl, {f"{P}S1": {
            "平时成绩": 80, "期中理论": 80, "期中实践": 90,
            "期末理论": 70, "期末实践": 90,
        }})
        body = upload(client, t1, "grades", data).json()
        check(body["success"] == 1 and body["failed"] == 0, "上传成绩", str(body))
        with SessionLocal() as db:
            g = db.execute(select(Grade).where(Grade.student_id == s1_id)).scalar_one()
            check(g.status == "DRAFT", "成绩落库为草稿（待教师核对后提交）", g.status)
            check(g.theory_total_score == Decimal("75.00")
                  and g.practice_total_score == Decimal("88.00")
                  and g.total_score == Decimal("81.50"),
                  "ScoreService 统一计算：理论75 / 实践88 / 总评81.50",
                  f"{g.theory_total_score}/{g.practice_total_score}/{g.total_score}")
            check(g.mid_score == Decimal("85.00") and g.final_score == Decimal("80.00"),
                  "合并值由后端回写（期中85 / 期末80）")
            grade_id = g.id

        # 非法成绩
        bad = build(hdr, [[f"{P}S1", "E2E学生甲", "班", f"{P}C1", "课", tc_code,
                           101, "", "", "", "", "", ""]])
        body = upload(client, t1, "grades", bad).json()
        check(body["failed"] == 1 and "超出范围" in body["errors"][0]["message"],
              "非法成绩（101）被拒绝", str(body["errors"][0]))

        # ---------- 6. 草稿删除 ----------
        print("\n[6] 成绩草稿删除")
        r = client.delete(f"{BASE}/grades/{grade_id}", headers=t1)
        check(r.status_code == 204, "删除 DRAFT → 204")
        with SessionLocal() as db:
            gone = db.get(Grade, grade_id)
            check(gone is None, "MySQL：草稿行已删除")

        # SUBMITTED / APPROVED 拒绝
        for status, label in (("SUBMITTED", "待审核"), ("APPROVED", "已通过")):
            with SessionLocal() as db:
                g = Grade(student_id=s1_id, teaching_class_id=tc_id, course_id=c48p.id,
                          semester="2026-2027-1", credits=Decimal("3.0"),
                          regular_score=Decimal("80"), status=status)
                db.add(g)
                db.commit()
                gid = g.id
            r = client.delete(f"{BASE}/grades/{gid}", headers=t1)
            check(r.status_code == 409 and "只有草稿可以删除" in r.json()["message"],
                  f"{label}（{status}）不可删除 → 409")
            with SessionLocal() as db:
                check(db.get(Grade, gid) is not None, f"{label}成绩未被删除")
                # uk_grades 限制 (student, teaching_class) 唯一，清理后再测下一个状态
                db.delete(db.get(Grade, gid))
                db.commit()

        # ---------- 7. 教师越权 ----------
        print("\n[7] 对象级权限")
        t2 = login(client, f"{P}T2", f"{P}T2")
        with SessionLocal() as db:
            g = Grade(student_id=s1_id, teaching_class_id=tc_id, course_id=c48p.id,
                      semester="2026-2027-1", credits=Decimal("3.0"),
                      regular_score=Decimal("80"), status="DRAFT")
            db.add(g)
            db.commit()
            other_id = g.id
        r = client.delete(f"{BASE}/grades/{other_id}", headers=t2)
        check(r.status_code == 403, "教师乙删除教师甲的草稿 → 403")
        bad = build(hdr, [[f"{P}S1", "", "", "", "", tc_code,
                           80, 80, 90, 70, 90, "", ""]])
        body = upload(client, t2, "grades", bad).json()
        check(body["failed"] == 1 and "只能导入本人授课教学班" in body["errors"][0]["message"],
              "教师乙导入教师甲的教学班 → 拒绝", str(body["errors"][0]))

    cleanup()
    passed = sum(1 for ok, _ in results if ok)
    print(f"\n{'=' * 60}\nE2E 结果：{passed}/{len(results)} 通过")
    for ok, label in results:
        if not ok:
            print(f"  {FAIL} {label}")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
