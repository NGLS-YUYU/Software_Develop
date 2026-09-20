"""FIX-005 前端真实浏览器 E2E（Chrome + CDP）。

在真实浏览器里点真实按钮：登录 → 批量导入页 → 下载模板 → 选择文件 → 开始导入
→ 读取结果卡片与错误明细；再以教师身份进入成绩录入页 → 下载模板 → 导入成绩
→ 删除草稿（含确认弹窗）。

前置：
    uvicorn app.main:app --port 8000          # 前端 vite 代理指向 8000
    npm run dev                                # http://127.0.0.1:5173

用法：
    .venv/bin/python scripts/e2e_fix005_browser.py
"""

import asyncio
import io
import json
import shutil
import subprocess
import sys
import tempfile
import time
from decimal import Decimal
from pathlib import Path

import httpx
import websockets
from openpyxl import Workbook, load_workbook

BACKEND = "http://127.0.0.1:8000/api/v1"
FRONTEND = "http://127.0.0.1:5173"
CDP = "http://127.0.0.1:9222"
P = "BR5"
DOWNLOAD_DIR = Path(tempfile.mkdtemp(prefix="uams_dl_"))

results: list[tuple[bool, str]] = []


def check(cond: bool, label: str, detail: str = "") -> None:
    results.append((bool(cond), label))
    print(f"  {'✅' if cond else '❌'} {label}{(' — ' + detail) if detail else ''}")


class Browser:
    """极简 CDP 客户端：只实现本脚本需要的命令。"""

    def __init__(self, ws):
        self.ws = ws
        self._id = 0

    async def send(self, method: str, **params):
        self._id += 1
        msg_id = self._id
        await self.ws.send(json.dumps({"id": msg_id, "method": method, "params": params}))
        while True:
            data = json.loads(await self.ws.recv())
            if data.get("id") == msg_id:
                if "error" in data:
                    raise RuntimeError(f"{method}: {data['error']}")
                return data.get("result", {})

    async def js(self, expr: str):
        r = await self.send(
            "Runtime.evaluate",
            expression=expr,
            returnByValue=True,
            awaitPromise=True,
        )
        res = r.get("result", {})
        if r.get("exceptionDetails"):
            raise RuntimeError(f"JS 异常：{r['exceptionDetails'].get('text')}")
        return res.get("value")

    async def goto(self, url: str):
        await self.send("Page.navigate", url=url)
        await self.wait_for("document.readyState === 'complete'", timeout=20)

    async def wait_for(self, expr: str, timeout: float = 15, interval: float = 0.2):
        deadline = time.time() + timeout
        while time.time() < deadline:
            try:
                if await self.js(expr):
                    return True
            except RuntimeError:
                pass
            await asyncio.sleep(interval)
        return False

    async def set_file(self, selector: str, path: Path):
        doc = await self.send("DOM.getDocument", depth=-1)
        node = await self.send(
            "DOM.querySelector", nodeId=doc["root"]["nodeId"], selector=selector
        )
        if not node.get("nodeId"):
            raise RuntimeError(f"找不到元素：{selector}")
        await self.send("DOM.setFileInputFiles", files=[str(path)], nodeId=node["nodeId"])

    async def click_text(self, selector: str, text: str) -> bool:
        return await self.js(
            f"""(() => {{
              const els = [...document.querySelectorAll({json.dumps(selector)})];
              const el = els.find(e => (e.textContent || '').includes({json.dumps(text)}));
              if (!el) return false;
              el.click(); return true;
            }})()"""
        )


FILL_INPUT = """
((sel, val) => {
  const el = document.querySelector(sel);
  if (!el) return false;
  const setter = Object.getOwnPropertyDescriptor(
    window.HTMLInputElement.prototype, 'value').set;
  setter.call(el, val);
  el.dispatchEvent(new Event('input', { bubbles: true }));
  el.dispatchEvent(new Event('change', { bubbles: true }));
  return true;
})(%s, %s)
"""


def teacher_file() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["工号", "姓名", "学院", "职称", "性别"])
    ws.append([f"{P}T1", "浏览器教师", "计算机科学与技术学院", "教授", "女"])
    ws.append([f"{P}T2", "浏览器教师B", "计算机科学与技术学院", "讲师", "男"])
    p = DOWNLOAD_DIR / "teachers.xlsx"
    wb.save(p)
    return p


def bad_teacher_file() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["工号", "姓名", "学院"])
    ws.append([f"{P}T3", "错误行", "不存在的学院"])
    p = DOWNLOAD_DIR / "teachers_bad.xlsx"
    wb.save(p)
    return p


def cleanup() -> None:
    from sqlalchemy import delete, select

    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from app.db.session import SessionLocal
    from app.models.course import Course
    from app.models.grade import Grade
    from app.models.person import Student, Teacher
    from app.models.role import UserRole
    from app.models.selection import CourseSelection
    from app.models.teaching import Schedule, TeachingClass, TeachingTask
    from app.models.user import User

    with SessionLocal() as db:
        users = select(User.id).where(User.username.like(f"{P}%"))
        courses = select(Course.id).where(Course.code.like(f"{P}%"))
        tcs = select(TeachingClass.id).where(TeachingClass.code.like(f"{P}%"))
        db.execute(delete(UserRole).where(UserRole.user_id.in_(users)))
        db.execute(delete(Grade).where(Grade.teaching_class_id.in_(tcs)))
        db.execute(delete(CourseSelection).where(CourseSelection.teaching_class_id.in_(tcs)))
        db.execute(delete(Schedule).where(Schedule.teaching_class_id.in_(tcs)))
        db.execute(delete(TeachingClass).where(TeachingClass.code.like(f"{P}%")))
        db.execute(delete(TeachingTask).where(TeachingTask.course_id.in_(courses)))
        db.execute(delete(Course).where(Course.code.like(f"{P}%")))
        db.execute(delete(Student).where(Student.student_no.like(f"{P}%")))
        db.execute(delete(Teacher).where(Teacher.teacher_no.like(f"{P}%")))
        db.execute(delete(User).where(User.username.like(f"{P}%")))
        db.commit()


async def main() -> int:
    cleanup()
    chrome = subprocess.Popen(
        [
            "google-chrome", "--headless=new", "--no-sandbox", "--disable-gpu",
            "--remote-debugging-port=9222",
            f"--user-data-dir={DOWNLOAD_DIR / 'profile'}",
            "about:blank",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        targets = None
        for _ in range(60):
            try:
                targets = httpx.get(f"{CDP}/json/list", timeout=2).json()
                if targets:
                    break
            except Exception:
                await asyncio.sleep(0.5)
        if not targets:
            print("❌ 无法连接 Chrome CDP")
            return 1

        page = next(t for t in targets if t["type"] == "page")
        async with websockets.connect(page["webSocketDebuggerUrl"], max_size=20_000_000) as ws:
            b = Browser(ws)
            await b.send("Page.enable")
            await b.send("DOM.enable")
            await b.send("Runtime.enable")
            await b.send("Page.setDownloadBehavior",
                         behavior="allow", downloadPath=str(DOWNLOAD_DIR))

            # ---------- 1. 登录（真实表单交互） ----------
            print("\n[1] 真实登录（管理员）")
            await b.goto(f"{FRONTEND}/login")
            check(await b.wait_for("document.querySelectorAll('input').length >= 2", timeout=30),
                  "登录页已渲染（Vue 已挂载）")
            await b.js(FILL_INPUT % ("'input:not([type=password])'", "'admin'"))
            await b.js(FILL_INPUT % ("'input[type=password]'", "'admin123'"))
            await b.click_text("button", "登 录")
            ok = await b.wait_for("location.pathname.includes('/admin')", timeout=30)
            check(ok, "表单登录成功并跳转教务端", await b.js("location.pathname"))

            # ---------- 2. 批量导入页 ----------
            print("\n[2] 批量导入页（真实按钮）")
            await b.goto(f"{FRONTEND}/admin/imports")
            check(await b.wait_for("!!document.querySelector('.el-tabs')", timeout=30),
                  "批量导入页已渲染")
            tabs = await b.js(
                "[...document.querySelectorAll('.el-tabs__item')].map(e=>e.textContent.trim())"
            )
            check(tabs == ["学生导入", "教师导入", "课程导入", "成绩导入"],
                  "四个导入 Tab 齐全", str(tabs))

            # 下载模板（真实浏览器下载）
            await b.click_text(".el-tabs__item", "教师导入")
            await asyncio.sleep(0.4)
            await b.click_text("button", "① 下载模板")
            await asyncio.sleep(2.5)
            downloaded = [f for f in DOWNLOAD_DIR.glob("*.xlsx") if "template" in f.name]
            check(bool(downloaded), "点击「下载模板」触发真实文件下载",
                  str([f.name for f in downloaded]))
            if downloaded:
                ws0 = load_workbook(downloaded[0]).worksheets[0]
                check([c.value for c in ws0[1]][:3] == ["工号", "姓名", "学院"],
                      "下载到的模板首行是中文字段名")
                check(load_workbook(downloaded[0]).sheetnames == ["导入模板", "填写说明"],
                      "下载到的模板含「填写说明」工作表")

            # 上传（真实文件选择 + 真实按钮）
            await b.set_file("input[type=file]", teacher_file())
            await asyncio.sleep(0.4)
            check(await b.click_text("button", "③ 开始导入"), "点击「开始导入」")
            ok = await b.wait_for(
                "!!document.querySelector('.el-statistic')", timeout=20
            )
            check(ok, "结果卡片已渲染")
            stats = await b.js(
                "[...document.querySelectorAll('.el-statistic')].map(e=>e.innerText.replace(/\\s+/g,' '))"
            )
            check(any("2" in s for s in stats), "结果卡片显示成功 2 条", str(stats))

            # 错误明细
            await b.set_file("input[type=file]", bad_teacher_file())
            await asyncio.sleep(0.4)
            await b.click_text("button", "③ 开始导入")
            ok = await b.wait_for(
                "document.body.innerText.includes('学院不存在')", timeout=20
            )
            check(ok, "错误明细显示行号与原因",
                  (await b.js("""(() => {
                      const t = [...document.querySelectorAll('.el-table__row')]
                        .map(r => r.innerText.replace(/\\s+/g,' ')).join(' | ');
                      return t;
                  })()""")) or "")

            # 数据库校验
            sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
            from sqlalchemy import select as _select

            from app.db.session import SessionLocal
            from app.models.person import Teacher
            with SessionLocal() as db:
                t = db.execute(
                    _select(Teacher).where(Teacher.teacher_no == f"{P}T1")
                ).scalar_one_or_none()
            check(t is not None and t.title == "PROFESSOR",
                  "MySQL：浏览器导入的教师已落库（职称=教授）")

            # ---------- 3. 教师端：成绩模板 / 导入 / 删除草稿 ----------
            print("\n[3] 教师成绩录入（真实按钮）")
            await b.js("localStorage.clear()")
            await b.goto(f"{FRONTEND}/login")
            await b.wait_for("document.querySelectorAll('input').length >= 2", timeout=30)
            await b.js(FILL_INPUT % ("'input:not([type=password])'", "'T001'"))
            await b.js(FILL_INPUT % ("'input[type=password]'", "'T001'"))
            await b.click_text("button", "登 录")
            check(await b.wait_for("location.pathname.includes('/teacher')", timeout=30),
                  "教师登录成功", await b.js("location.pathname"))

            tok = httpx.post(f"{BACKEND}/auth/login",
                             json={"username": "T001", "password": "T001"},
                             timeout=20).json()["access_token"]

            # 为教师端 E2E 造一个隔离的教学班（不污染演示数据），归 T001 所有
            sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
            from sqlalchemy import select as _sel

            from app.db.session import SessionLocal
            from app.models.course import Course
            from app.models.grade import Grade
            from app.models.person import Student, Teacher
            from app.models.selection import CourseSelection, SelectionState
            from app.models.teaching import TeachingClass, TeachingTask
            from app.models.term import AcademicTerm

            with SessionLocal() as db:
                t1 = db.execute(
                    _sel(Teacher).where(Teacher.teacher_no == "T001")
                ).scalar_one()
                college_id = t1.college_id
                course = Course(
                    code=f"{P}C1", name="浏览器E2E课", college_id=college_id,
                    course_type="REQUIRED", credits=Decimal("3.0"),
                    total_hours=48, theory_hours=40, practice_hours=8,
                )
                db.add(course)
                db.flush()
                task = TeachingTask(semester="2026-2027-1", course_id=course.id,
                                    college_id=college_id, status="CONFIRMED")
                db.add(task)
                db.flush()
                term = db.execute(_sel(AcademicTerm).where(
                    AcademicTerm.academic_year == "2026-2027",
                    AcademicTerm.semester == 1,
                )).scalar_one()
                tc = TeachingClass(
                    code=f"{P}TC1", task_id=task.id, course_id=course.id,
                    semester="2026-2027-1", teacher_id=t1.id, capacity=60,
                    term_id=term.id, selection_status="OPEN",
                )
                db.add(tc)
                db.flush()
                students = list(db.execute(
                    _sel(Student).where(Student.student_no.in_(
                        ["20260001", "20260002", "20260003"]
                    )).order_by(Student.student_no)
                ).scalars())
                for s in students:
                    db.add(CourseSelection(
                        student_id=s.id, teaching_class_id=tc.id,
                        course_id=course.id, semester="2026-2027-1",
                        status=SelectionState.SELECTED,
                    ))
                db.commit()
                target = {"id": tc.id, "course_name": course.name, "code": tc.code}
                expected_nos = [s.student_no for s in students]
            check(len(expected_nos) == 3, "已为教师端 E2E 准备隔离教学班（3 名学生）",
                  str(expected_nos))

            await b.goto(f"{FRONTEND}/teacher/grades/{target['id']}")
            check(await b.wait_for(
                "document.body.innerText.includes('生成成绩单')", timeout=30),
                "成绩录入页已渲染（未生成成绩单态）")
            await b.click_text("button", "生成成绩单")
            check(await b.wait_for("!!document.querySelector('.el-table__row')", timeout=30),
                  "点击「生成成绩单」后表格渲染")
            before = await b.js("document.querySelectorAll('.el-table__row').length")
            check(before == 3, f"表格载入 {before} 名学生（成绩单已生成）")

            # 下载成绩模板（应带出学生名单 + 按考核结构裁剪列）
            await b.click_text("button", "下载导入模板")
            await asyncio.sleep(2.5)
            gfiles = sorted(DOWNLOAD_DIR.glob("grade_import_template_*.xlsx"))
            check(bool(gfiles), "点击「下载导入模板」触发真实下载",
                  str([f.name for f in gfiles]))
            if gfiles:
                wb = load_workbook(gfiles[-1])
                ws0 = wb.worksheets[0]
                hdr = [c.value for c in ws0[1]]
                check(hdr[:6] == ["学号", "姓名", "班级", "课程代码", "课程名称", "教学班"],
                      "成绩模板固定列由系统生成", str(hdr[:6]))
                check(hdr[6:] == ["平时成绩", "期中理论", "期中实践", "期末理论",
                                  "期末实践", "理论补考", "实践补考"],
                      "48h+实践 → 理论/实践拆分列（无合并「期中成绩」列）", str(hdr[6:]))
                roster = [
                    str(ws0.cell(r, 1).value).strip()
                    for r in range(2, ws0.max_row + 1)
                    if ws0.cell(r, 1).value
                ]
                check(roster == expected_nos, f"模板自动带出学生名单 {roster}")

                # 只填「原始成绩」列（补考列留给不及格场景，见下一步）
                makeup_cols = {
                    i for i, label in enumerate(hdr, start=1)
                    if "补考" in str(label)
                }
                fill_cols = [
                    i for i in range(7, len(hdr) + 1) if i not in makeup_cols
                ]
                for r in range(2, ws0.max_row + 1):
                    for ci in fill_cols:
                        ws0.cell(r, ci).value = 85
                filled_path = DOWNLOAD_DIR / "grade_filled.xlsx"
                wb.save(filled_path)

                await b.set_file("#app input[type=file]", filled_path)
                await asyncio.sleep(0.4)
                await b.click_text("button", "导入成绩")
                ok = await b.wait_for(
                    "document.body.innerText.includes('导入结果')", timeout=25
                )
                check(ok, "点击「导入成绩」后页面显示导入结果")
                await asyncio.sleep(1.5)

                with SessionLocal() as db:
                    rows = list(db.execute(
                        _sel(Grade).where(Grade.teaching_class_id == target["id"])
                    ).scalars())
                check(len(rows) == 3 and all(g.status == "DRAFT" for g in rows),
                      "MySQL：浏览器导入的 3 条成绩落库为草稿",
                      str([(g.status, str(g.total_score)) for g in rows]))
                # 48h+实践，全部 85 分 → 理论=实践=总评 85.00
                check(all(g.total_score == Decimal("85.00") for g in rows),
                      "ScoreService 计算：全 85 分 → 总评 85.00",
                      str([str(g.total_score) for g in rows]))

                # 补考列：期末已及格仍填补考 → 后端逐行拒绝（错误明细可见）
                first_makeup_col = min(makeup_cols)
                wb2 = load_workbook(filled_path)
                ws2 = wb2.worksheets[0]
                for r in range(2, ws2.max_row + 1):
                    ws2.cell(r, first_makeup_col).value = 90
                makeup_path = DOWNLOAD_DIR / "grade_makeup.xlsx"
                wb2.save(makeup_path)
                await b.set_file("#app input[type=file]", makeup_path)
                await asyncio.sleep(0.4)
                await b.click_text("button", "导入成绩")
                ok = await b.wait_for(
                    "document.body.innerText.includes('不及格')", timeout=25
                )
                check(ok, "期末已及格却填补考 → 页面显示逐行拒绝原因",
                      (await b.js("""(() => {
                          const row = document.querySelector('.el-table__row');
                          return row ? row.innerText.replace(/\\s+/g,' ').slice(0,60) : '';
                      })()""")) or "")
                with SessionLocal() as db:
                    untouched = list(db.execute(
                        _sel(Grade).where(Grade.teaching_class_id == target["id"])
                    ).scalars())
                check(all(g.total_score == Decimal("85.00") for g in untouched),
                      "被拒绝的行未污染已有草稿")

                # 删除草稿（真实按钮 + 二次确认）
                first_no = await b.js(
                    "document.querySelector('.el-table__row .el-table__cell:nth-child(2)').innerText.trim()"
                )
                clicked = await b.js("""(() => {
                    const row = document.querySelector('.el-table__row');
                    const btn = [...row.querySelectorAll('button')]
                      .find(e => e.textContent.includes('删除草稿'));
                    if (!btn) return false;
                    btn.click(); return true;
                })()""")
                check(clicked, "点击行内「删除草稿」")
                check(await b.wait_for("!!document.querySelector('.el-message-box')", timeout=10),
                      "弹出二次确认对话框",
                      (await b.js("document.querySelector('.el-message-box')?.innerText.slice(0,60)")) or "")
                await b.js("""(() => {
                    const btn = [...document.querySelectorAll('.el-message-box__btns button')]
                      .find(e => e.textContent.includes('删除'));
                    btn && btn.click();
                })()""")
                await asyncio.sleep(2)
                check(await b.js("document.body.innerText.includes('删除成功')"),
                      "提示「删除成功」并刷新列表")
                with SessionLocal() as db:
                    stu = db.execute(
                        _sel(Student).where(Student.student_no == first_no.strip())
                    ).scalar_one_or_none()
                    left = db.execute(
                        _sel(Grade).where(
                            Grade.teaching_class_id == target["id"],
                            Grade.student_id == (stu.id if stu else -1),
                        )
                    ).scalar_one_or_none()
                check(left is None, f"MySQL：{first_no.strip()} 的草稿已删除")
                check(await b.js("document.querySelectorAll('.el-table__row').length") == 2,
                      "表格刷新为 2 行（该生从成绩单移除）")
                # 名单接口确认该生回到「未录入」，且可重新生成成绩单
                after = httpx.get(
                    f"{BACKEND}/grades/teaching-classes/{target['id']}/students",
                    headers={"Authorization": f"Bearer {tok}"}, timeout=20,
                ).json()
                check(
                    any(r["student_no"] == first_no.strip()
                        and r["entry_status"] == "NOT_ENTERED" for r in after),
                    f"名单接口确认 {first_no.strip()} 回到「未录入」",
                )

            print(f"\n{'=' * 60}\n浏览器 E2E：{sum(1 for ok, _ in results if ok)}/{len(results)} 通过")
            for ok, label in results:
                if not ok:
                    print(f"  ❌ {label}")
            return 0 if all(ok for ok, _ in results) else 1
    finally:
        chrome.terminate()
        shutil.rmtree(DOWNLOAD_DIR, ignore_errors=True)
        cleanup()


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
