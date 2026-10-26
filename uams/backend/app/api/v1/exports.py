"""业务数据 Excel 导出（F-17）。"""
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.api.deps import RequirePermissions
from app.db.session import get_db
from app.models.course import Course
from app.models.grade import Grade, GradeStatus
from app.models.person import Student
from app.models.selection import CourseSelection, SelectionState

router = APIRouter(prefix="/exports", tags=["导出"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(wb: Workbook, filename: str) -> Response:
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _name() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


@router.get("/students", summary="导出学生名单")
def export_students(
    db: Session = Depends(get_db),
    _=Depends(RequirePermissions("student:export")),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["学号", "姓名", "学院", "专业", "行政班", "年级", "学籍状态", "邮箱", "手机"])
    rows = db.execute(
        select(Student).options(
            selectinload(Student.user), selectinload(Student.klass),
            selectinload(Student.major), selectinload(Student.college),
        ).order_by(Student.student_no)
    ).scalars()
    for s in rows:
        ws.append([
            s.student_no, s.user.real_name if s.user else "",
            s.college.name if s.college else "", s.major.name if s.major else "",
            s.klass.name if s.klass else "", s.enrollment_year,
            s.academic_status, s.user.email if s.user else "",
            s.user.phone if s.user else "",
        ])
    return _xlsx(wb, f"students_{_name()}.xlsx")


@router.get("/courses", summary="导出课程库")
def export_courses(
    db: Session = Depends(get_db),
    _=Depends(RequirePermissions("course:read")),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "课程库"
    ws.append(["课程代码", "课程名称", "开课学院", "课程性质", "学分", "总学时", "理论学时", "实践学时", "状态"])
    rows = db.execute(
        select(Course).options(selectinload(Course.college)).order_by(Course.code)
    ).scalars()
    for c in rows:
        ws.append([c.code, c.name, c.college.name if c.college else "", c.course_type,
                   float(c.credits), c.total_hours, c.theory_hours, c.practice_hours, c.status])
    return _xlsx(wb, f"courses_{_name()}.xlsx")


@router.get("/grades", summary="导出已审核成绩")
def export_grades(
    semester: str | None = Query(None),
    db: Session = Depends(get_db),
    _=Depends(RequirePermissions("grade:export")),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "已审核成绩"
    ws.append(["学号", "姓名", "行政班", "课程代码", "课程名称", "教学班", "学期",
               "平时", "期中", "期末", "总评", "绩点", "学分", "是否通过", "审核时间"])
    stmt = select(Grade).where(Grade.status == GradeStatus.APPROVED).options(
        selectinload(Grade.student).selectinload(Student.user),
        selectinload(Grade.student).selectinload(Student.klass),
        selectinload(Grade.course), selectinload(Grade.teaching_class),
    )
    if semester:
        stmt = stmt.where(Grade.semester == semester)
    for g in db.execute(stmt.order_by(Grade.semester, Grade.student_id)).scalars():
        s = g.student
        ws.append([
            s.student_no if s else "", s.user.real_name if s and s.user else "",
            s.klass.name if s and s.klass else "",
            g.course.code if g.course else "", g.course.name if g.course else "",
            g.teaching_class.code if g.teaching_class else "", g.semester,
            float(g.regular_score) if g.regular_score is not None else "",
            float(g.mid_score) if g.mid_score is not None else "",
            float(g.final_score) if g.final_score is not None else "",
            float(g.total_score) if g.total_score is not None else "",
            float(g.grade_point) if g.grade_point is not None else "",
            float(g.credits), "是" if g.is_passed else "否",
            g.approved_at.strftime("%Y-%m-%d %H:%M") if g.approved_at else "",
        ])
    return _xlsx(wb, f"grades_{_name()}.xlsx")


@router.get("/selections", summary="导出选课结果")
def export_selections(
    semester: str | None = Query(None),
    db: Session = Depends(get_db),
    _=Depends(RequirePermissions("selection:manage")),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "选课结果"
    ws.append(["学号", "姓名", "行政班", "课程代码", "课程名称", "教学班", "学期", "来源", "选课时间"])
    stmt = select(CourseSelection).where(
        CourseSelection.status == SelectionState.SELECTED
    ).options(
        selectinload(CourseSelection.student).selectinload(Student.user),
        selectinload(CourseSelection.student).selectinload(Student.klass),
        selectinload(CourseSelection.course), selectinload(CourseSelection.teaching_class),
    )
    if semester:
        stmt = stmt.where(CourseSelection.semester == semester)
    for cs in db.execute(stmt.order_by(CourseSelection.semester, CourseSelection.student_id)).scalars():
        s = cs.student
        ws.append([
            s.student_no if s else "", s.user.real_name if s and s.user else "",
            s.klass.name if s and s.klass else "",
            cs.course.code if cs.course else "", cs.course.name if cs.course else "",
            cs.teaching_class.code if cs.teaching_class else "", cs.semester,
            cs.source or "", cs.selected_at.strftime("%Y-%m-%d %H:%M") if cs.selected_at else "",
        ])
    return _xlsx(wb, f"selections_{_name()}.xlsx")
